"""Import dei file Excel: lettura, mappatura colonne, controllo e scrittura nelle tabelle."""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Cliente, PassaggioOfficina, Sede, Veicolo

IGNORA = "ignora"
EXTRA = "extra"

# Campi del CRM a cui si può abbinare una colonna del file.
CAMPI: dict[str, str] = {
    "cliente.codice": "Codice cliente",
    "cliente.cognome": "Cognome / Ragione sociale",
    "cliente.nome": "Nome",
    "cliente.tipo": "Privato o azienda",
    "cliente.telefono_cellulare": "Telefono cellulare",
    "cliente.telefono_fisso": "Telefono fisso",
    "cliente.email": "Email",
    "veicolo.targa": "Targa",
    "veicolo.telaio": "Telaio",
    "veicolo.marca": "Marca",
    "veicolo.modello": "Modello",
    "veicolo.tipologia": "Tipologia veicolo",
    "veicolo.data_immatricolazione": "Data immatricolazione",
    "passaggio.numero_or": "Numero O.R. / pratica",
    "passaggio.descrizione": "Descrizione intervento",
    "passaggio.sede": "Sede officina",
    "passaggio.data_apertura": "Data apertura O.R.",
    "passaggio.data_chiusura": "Data chiusura O.R.",
}

CAMPI_DATA = {
    "veicolo.data_immatricolazione",
    "passaggio.data_apertura",
    "passaggio.data_chiusura",
}

# Intestazioni note (normalizzate) -> campo. Copre il file di esempio REGIE AUTO.
SINONIMI: dict[str, str] = {
    "cod cliente": "cliente.codice",
    "codice cliente": "cliente.codice",
    "cognome": "cliente.cognome",
    "ragione sociale": "cliente.cognome",
    "nominativo": "cliente.cognome",
    "nome": "cliente.nome",
    "tel cellulare": "cliente.telefono_cellulare",
    "cellulare": "cliente.telefono_cellulare",
    "tel domicilio": "cliente.telefono_fisso",
    "telefono": "cliente.telefono_fisso",
    "tel fisso": "cliente.telefono_fisso",
    "e mail": "cliente.email",
    "email": "cliente.email",
    "targa": "veicolo.targa",
    "telaio": "veicolo.telaio",
    "vin": "veicolo.telaio",
    "marca": "veicolo.marca",
    "modello": "veicolo.modello",
    "tipologia": "veicolo.tipologia",
    "tipologia veicolo": "veicolo.tipologia",
    "immatricolazione": "veicolo.data_immatricolazione",
    "data immatricolazione": "veicolo.data_immatricolazione",
    "n o r pratica": "passaggio.numero_or",
    "numero or": "passaggio.numero_or",
    "intervento richiesta cliente": "passaggio.descrizione",
    "servizio": "passaggio.sede",
    "sede": "passaggio.sede",
    "data creaz": "passaggio.data_apertura",
    "data emis": "passaggio.data_chiusura",
    # Oggi compilate a mano dalle operatrici: nel CRM diventano il registro chiamate.
    "data chiamata": IGNORA,
    "caller": IGNORA,
    "app preso": IGNORA,
}


class FileNonValido(ValueError):
    pass


# ---------------------------------------------------------------- lettura file


def _cella(v):
    """Rende il valore di una cella salvabile in JSON."""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def leggi_file(contenuto: bytes) -> tuple[list[str], list[tuple[int, list]]]:
    """Legge .xlsx, .xls o .csv e restituisce intestazioni e righe non vuote."""
    if contenuto[:2] == b"PK":
        righe = _leggi_xlsx(contenuto)
    elif contenuto[:4] == b"\xd0\xcf\x11\xe0":
        righe = _leggi_xls(contenuto)
    else:
        righe = _leggi_csv(contenuto)

    # Ogni riga tiene il suo numero nel file, per indicare dove sono gli errori.
    numerate = [(i + 1, [_cella(v) for v in r]) for i, r in enumerate(righe)]
    numerate = [(i, r) for i, r in numerate if any(v is not None for v in r)]
    if not numerate:
        raise FileNonValido("Il file è vuoto")

    intestazioni = [str(v) if v is not None else "" for v in numerate[0][1]]
    corpo = numerate[1:]
    # Toglie le colonne finali senza intestazione e senza dati.
    while intestazioni and not intestazioni[-1] and all(
        len(r) < len(intestazioni) or r[len(intestazioni) - 1] is None for _, r in corpo
    ):
        intestazioni.pop()
    n = len(intestazioni)
    if n == 0:
        raise FileNonValido("Il file non ha intestazioni nella prima riga")
    intestazioni = [h or f"Colonna {i + 1}" for i, h in enumerate(intestazioni)]
    return intestazioni, [(i, (r + [None] * n)[:n]) for i, r in corpo]


def _leggi_xlsx(contenuto: bytes) -> list[list]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(contenuto), read_only=True, data_only=True)
    except Exception as exc:
        raise FileNonValido("Il file Excel non si riesce a leggere") from exc
    ws = wb.worksheets[0]
    righe = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return righe


def _leggi_xls(contenuto: bytes) -> list[list]:
    import xlrd

    try:
        wb = xlrd.open_workbook(file_contents=contenuto)
    except Exception as exc:
        raise FileNonValido("Il file Excel .xls non si riesce a leggere") from exc
    ws = wb.sheet_by_index(0)
    righe = []
    for i in range(ws.nrows):
        riga = []
        for c in ws.row(i):
            if c.ctype == xlrd.XL_CELL_DATE:
                riga.append(xlrd.xldate.xldate_as_datetime(c.value, wb.datemode))
            elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                riga.append(None)
            else:
                riga.append(c.value)
        righe.append(riga)
    return righe


def _leggi_csv(contenuto: bytes) -> list[list]:
    for codifica in ("utf-8-sig", "cp1252"):
        try:
            testo = contenuto.decode(codifica)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileNonValido("Formato non riconosciuto: usa un file .xlsx, .xls o .csv")
    if "\x00" in testo:
        raise FileNonValido("Formato non riconosciuto: usa un file .xlsx, .xls o .csv")
    try:
        dialetto = csv.Sniffer().sniff(testo[:4096], delimiters=";,\t")
    except csv.Error:

        class dialetto(csv.excel):
            delimiter = ";"

    return [list(r) for r in csv.reader(io.StringIO(testo), dialetto)]


# ------------------------------------------------------------------ mappatura


def _norm_intestazione(h: str) -> str:
    h = unicodedata.normalize("NFKD", h).encode("ascii", "ignore").decode().lower()
    return " ".join(re.findall(r"[a-z0-9]+", h))


def suggerisci_mappatura(intestazioni: list[str], salvata: dict | None = None) -> list[str]:
    """Una voce per colonna: campo CRM, 'extra' (si conserva) o 'ignora'."""
    usati: set[str] = set()
    risultato = []
    for h in intestazioni:
        campo = None
        if salvata and h in salvata:
            campo = salvata[h]
        else:
            campo = SINONIMI.get(_norm_intestazione(h))
        if campo in CAMPI and campo in usati:
            campo = None
        if campo is None:
            campo = EXTRA
        if campo in CAMPI:
            usati.add(campo)
        risultato.append(campo)
    return risultato


def controlla_mappatura(intestazioni: list[str], mappatura: list[str]) -> list[str]:
    """Errori bloccanti nella mappatura scelta (vuoto = va bene)."""
    errori = []
    if len(mappatura) != len(intestazioni):
        return ["La mappatura non corrisponde alle colonne del file"]
    visti: dict[str, str] = {}
    for h, campo in zip(intestazioni, mappatura, strict=True):
        if campo not in CAMPI and campo not in (EXTRA, IGNORA):
            errori.append(f"Campo sconosciuto per la colonna «{h}»")
        elif campo in CAMPI:
            if campo in visti:
                errori.append(
                    f"«{CAMPI[campo]}» è abbinato a due colonne: «{visti[campo]}» e «{h}»"
                )
            visti[campo] = h
    if "cliente.codice" not in visti and "cliente.cognome" not in visti:
        errori.append("Serve almeno una colonna per il cliente: codice cliente o cognome")
    if "veicolo.targa" not in visti and "veicolo.telaio" not in visti:
        errori.append("Serve almeno una colonna per il veicolo: targa o telaio")
    return errori


# --------------------------------------------------------- controllo dei valori


@dataclass
class RigaControllata:
    numero: int
    valori: dict = field(default_factory=dict)
    extra: dict = field(default_factory=dict)
    errori: list[str] = field(default_factory=list)  # la riga viene scartata
    avvisi: list[str] = field(default_factory=list)  # la riga entra, con un dato corretto o vuoto

    @property
    def valida(self) -> bool:
        return not self.errori


def _testo(v) -> str | None:
    if v is None:
        return None
    s = " ".join(str(v).split())
    return s or None


def _data(v) -> date | None:
    """Accetta date Excel, testo gg/mm/aaaa o aaaa-mm-gg e numeri seriali Excel."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 <= v <= 80000:
            return date(1899, 12, 30) + timedelta(days=int(v))
        raise ValueError
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError


def _telefono(v) -> tuple[str | None, str | None]:
    """Restituisce (numero pulito, avviso). Il numero è None se non è valido."""
    if v is None or v == "":
        return None, None
    grezzo = str(v).strip()
    cifre = re.sub(r"\D", "", grezzo)
    if grezzo.startswith("+") and not cifre.startswith("39"):
        # Numero estero: si tiene così com'è.
        return "+" + cifre, None
    if cifre.startswith("0039"):
        cifre = cifre[4:]
    elif cifre.startswith("39") and len(cifre) >= 11:
        cifre = cifre[2:]
    if cifre.startswith("3") and 9 <= len(cifre) <= 10:
        return cifre, None
    if cifre.startswith("0") and 6 <= len(cifre) <= 11:
        return cifre, None
    if 5 <= len(cifre) <= 10 and not cifre.startswith(("0", "3")):
        # Excel ha salvato il fisso come numero e ha perso lo zero iniziale.
        return "0" + cifre, f"aggiunto lo 0 iniziale al numero {grezzo}"
    return None, f"numero di telefono non valido: {grezzo}"


def _tipo_cliente(v) -> str | None:
    s = (_testo(v) or "").lower()
    if not s:
        return None
    if s.startswith(("p", "f")):  # privato, persona fisica
        return "privato"
    if s.startswith(("a", "s", "g")):  # azienda, società, giuridica
        return "azienda"
    return None


def controlla_riga(numero: int, dati: list, intestazioni: list[str], mappatura: list[str]) -> RigaControllata:
    r = RigaControllata(numero=numero)
    grezzi: dict[str, object] = {}
    for h, campo, v in zip(intestazioni, mappatura, dati, strict=True):
        if campo == IGNORA:
            continue
        if campo == EXTRA:
            if v is not None:
                r.extra[h] = v
            continue
        grezzi[campo] = v

    for campo, v in grezzi.items():
        etichetta = CAMPI[campo]
        if campo in CAMPI_DATA:
            try:
                r.valori[campo] = _data(v)
            except ValueError:
                r.valori[campo] = None
                r.avvisi.append(f"{etichetta}: data non valida ({v})")
        elif campo in ("cliente.telefono_cellulare", "cliente.telefono_fisso"):
            numero_tel, avviso = _telefono(v)
            r.valori[campo] = numero_tel
            if avviso:
                r.avvisi.append(f"{etichetta}: {avviso}")
        elif campo == "cliente.email":
            email = (_testo(v) or "").lower() or None
            if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
                r.avvisi.append(f"Email non valida: {v}")
                email = None
            r.valori[campo] = email
        elif campo == "cliente.tipo":
            r.valori[campo] = _tipo_cliente(v)
        elif campo == "veicolo.targa":
            t = re.sub(r"[\s\-.]", "", str(v)).upper() if v is not None else ""
            r.valori[campo] = t or None
        elif campo == "veicolo.telaio":
            t = re.sub(r"\s", "", str(v)).upper() if v is not None else ""
            r.valori[campo] = t or None
        elif campo in ("veicolo.marca",):
            t = _testo(v)
            r.valori[campo] = t.upper() if t else None
        else:
            r.valori[campo] = _testo(v)

    nominativo = " ".join(
        p for p in (r.valori.get("cliente.cognome"), r.valori.get("cliente.nome")) if p
    )
    r.valori["cliente.nominativo"] = nominativo or None

    if not r.valori.get("cliente.codice") and not nominativo:
        r.errori.append("manca il cliente (né codice né nominativo)")
    if not r.valori.get("veicolo.targa") and not r.valori.get("veicolo.telaio"):
        r.errori.append("manca il veicolo (né targa né telaio)")
    if (
        "cliente.telefono_cellulare" in grezzi or "cliente.telefono_fisso" in grezzi
    ) and not (r.valori.get("cliente.telefono_cellulare") or r.valori.get("cliente.telefono_fisso")):
        r.avvisi.append("nessun numero di telefono")
    return r


def anteprima(intestazioni: list[str], righe: list[tuple[int, list]], mappatura: list[str], max_righe: int = 50) -> dict:
    controllate = [controlla_riga(n, d, intestazioni, mappatura) for n, d in righe]
    codici_cliente = {
        r.valori.get("cliente.codice") or r.valori.get("cliente.nominativo")
        for r in controllate
        if r.valida
    }
    veicoli = {
        r.valori.get("veicolo.telaio") or r.valori.get("veicolo.targa")
        for r in controllate
        if r.valida
    }
    problemi = [
        {"riga": r.numero, "tipo": "errore", "messaggio": m}
        for r in controllate
        for m in r.errori
    ] + [
        {"riga": r.numero, "tipo": "avviso", "messaggio": m}
        for r in controllate
        for m in r.avvisi
    ]
    problemi.sort(key=lambda p: (p["riga"], p["tipo"]))
    return {
        "righe_totali": len(controllate),
        "righe_valide": sum(1 for r in controllate if r.valida),
        "righe_scartate": sum(1 for r in controllate if not r.valida),
        "righe_con_avvisi": sum(1 for r in controllate if r.valida and r.avvisi),
        "clienti_distinti": len(codici_cliente),
        "veicoli_distinti": len(veicoli),
        "problemi": problemi[:500],
        "esempio": [
            {
                "riga": r.numero,
                "valori": {k: (v.isoformat() if isinstance(v, date) else v) for k, v in r.valori.items()},
                "errori": r.errori,
                "avvisi": r.avvisi,
            }
            for r in controllate[:max_righe]
        ],
    }


# ------------------------------------------------------------------- scrittura


@dataclass
class Esito:
    clienti_nuovi: int = 0
    clienti_aggiornati: int = 0
    veicoli_nuovi: int = 0
    veicoli_aggiornati: int = 0
    passaggi_nuovi: int = 0
    righe_scartate: int = 0
    problemi: list[dict] = field(default_factory=list)


def _aggiorna(obj, valori: dict) -> bool:
    """Scrive solo i valori non vuoti: una cella vuota nel file non cancella un dato già presente."""
    cambiato = False
    for attr, v in valori.items():
        if v is not None and getattr(obj, attr) != v:
            setattr(obj, attr, v)
            cambiato = True
    return cambiato


class Scrittore:
    def __init__(self, db: Session):
        self.db = db
        self.esito = Esito()
        self._clienti_nuovi: set[int] = set()
        self._clienti_aggiornati: set[int] = set()
        self._veicoli_nuovi: set[int] = set()
        self._veicoli_aggiornati: set[int] = set()
        self._sedi: dict[str, Sede] = {}

    def _sede(self, nome: str | None) -> Sede | None:
        if not nome:
            return None
        if nome not in self._sedi:
            sede = self.db.scalar(select(Sede).where(Sede.nome == nome))
            if sede is None:
                sede = Sede(nome=nome)
                self.db.add(sede)
                self.db.flush()
            self._sedi[nome] = sede
        return self._sedi[nome]

    def _cliente(self, v: dict) -> Cliente:
        codice = v.get("cliente.codice")
        cliente = None
        if codice:
            cliente = self.db.scalar(select(Cliente).where(Cliente.codice == codice))
        else:
            # Senza codice: stesso nominativo e stesso telefono.
            for campo, col in (
                ("cliente.telefono_cellulare", Cliente.telefono_cellulare),
                ("cliente.telefono_fisso", Cliente.telefono_fisso),
            ):
                if v.get(campo) and cliente is None:
                    cliente = self.db.scalar(
                        select(Cliente).where(
                            col == v[campo], Cliente.nominativo == v.get("cliente.nominativo")
                        ).limit(1)
                    )
        dati = {
            "codice": codice,
            "nominativo": v.get("cliente.nominativo"),
            "tipo": v.get("cliente.tipo"),
            "telefono_cellulare": v.get("cliente.telefono_cellulare"),
            "telefono_fisso": v.get("cliente.telefono_fisso"),
            "email": v.get("cliente.email"),
        }
        if cliente is None:
            cliente = Cliente(extra={})
            dati["nominativo"] = dati["nominativo"] or f"Cliente {codice}"
            _aggiorna(cliente, dati)
            self.db.add(cliente)
            self.db.flush()
            self._clienti_nuovi.add(cliente.id)
        elif _aggiorna(cliente, dati) and cliente.id not in self._clienti_nuovi:
            self._clienti_aggiornati.add(cliente.id)
        return cliente

    def _veicolo(self, v: dict, cliente: Cliente, extra: dict) -> Veicolo:
        telaio, targa = v.get("veicolo.telaio"), v.get("veicolo.targa")
        veicolo = None
        if telaio:
            veicolo = self.db.scalar(select(Veicolo).where(Veicolo.telaio == telaio))
        if veicolo is None and targa:
            veicolo = self.db.scalar(
                select(Veicolo).where(Veicolo.targa == targa).order_by(Veicolo.id.desc()).limit(1)
            )
            # Stessa targa ma telaio diverso: è un altro veicolo (targa riassegnata).
            if veicolo is not None and telaio and veicolo.telaio and veicolo.telaio != telaio:
                veicolo = None
        dati = {
            "cliente_id": cliente.id,
            "targa": targa,
            "telaio": telaio,
            "marca": v.get("veicolo.marca"),
            "modello": v.get("veicolo.modello"),
            "tipologia": v.get("veicolo.tipologia"),
            "data_immatricolazione": v.get("veicolo.data_immatricolazione"),
        }
        if veicolo is None:
            veicolo = Veicolo(extra=extra)
            _aggiorna(veicolo, dati)
            self.db.add(veicolo)
            self.db.flush()
            self._veicoli_nuovi.add(veicolo.id)
        else:
            cambiato = _aggiorna(veicolo, dati)
            if extra and {**(veicolo.extra or {}), **extra} != veicolo.extra:
                veicolo.extra = {**(veicolo.extra or {}), **extra}
                cambiato = True
            if cambiato and veicolo.id not in self._veicoli_nuovi:
                self._veicoli_aggiornati.add(veicolo.id)
        return veicolo

    def _passaggio(self, v: dict, veicolo: Veicolo) -> None:
        dati = {
            "numero_or": v.get("passaggio.numero_or"),
            "descrizione": v.get("passaggio.descrizione"),
            "data_apertura": v.get("passaggio.data_apertura"),
            "data_chiusura": v.get("passaggio.data_chiusura"),
        }
        if not any(dati.values()):
            return
        sede = self._sede(v.get("passaggio.sede"))
        passaggio = None
        if dati["numero_or"]:
            passaggio = self.db.scalar(
                select(PassaggioOfficina).where(PassaggioOfficina.numero_or == dati["numero_or"])
            )
        else:
            passaggio = self.db.scalar(
                select(PassaggioOfficina).where(
                    PassaggioOfficina.veicolo_id == veicolo.id,
                    PassaggioOfficina.data_apertura == dati["data_apertura"],
                    PassaggioOfficina.descrizione == dati["descrizione"],
                )
            )
        if passaggio is None:
            passaggio = PassaggioOfficina(veicolo_id=veicolo.id)
            self.db.add(passaggio)
            self.esito.passaggi_nuovi += 1
        passaggio.veicolo_id = veicolo.id
        _aggiorna(passaggio, dati)
        if sede is not None:
            passaggio.sede_id = sede.id

        data = dati["data_chiusura"] or dati["data_apertura"]
        if data and (veicolo.data_ultimo_passaggio is None or data > veicolo.data_ultimo_passaggio):
            veicolo.data_ultimo_passaggio = data

    def scrivi(self, riga: RigaControllata) -> None:
        if not riga.valida:
            self.esito.righe_scartate += 1
            self.esito.problemi += [
                {"riga": riga.numero, "tipo": "errore", "messaggio": m} for m in riga.errori
            ]
            return
        self.esito.problemi += [
            {"riga": riga.numero, "tipo": "avviso", "messaggio": m} for m in riga.avvisi
        ]
        cliente = self._cliente(riga.valori)
        veicolo = self._veicolo(riga.valori, cliente, riga.extra)
        self._passaggio(riga.valori, veicolo)

    def chiudi(self) -> Esito:
        self.db.flush()
        self.esito.clienti_nuovi = len(self._clienti_nuovi)
        self.esito.clienti_aggiornati = len(self._clienti_aggiornati)
        self.esito.veicoli_nuovi = len(self._veicoli_nuovi)
        self.esito.veicoli_aggiornati = len(self._veicoli_aggiornati)
        return self.esito
